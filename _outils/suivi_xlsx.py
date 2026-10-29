#!/usr/bin/env python3
"""Génère _outils/suivi-candidatures-yasin-boukraa.xlsx à partir de _outils/suivi.json.
Usage : python3 _outils/suivi_xlsx.py   (puis recalc + upload Drive, voir METHODE.md)"""
import json, os, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
d = json.load(open(os.path.join(ROOT, "_outils", "suivi.json"), encoding="utf-8"))
OUT = os.path.join(ROOT, "_outils", "suivi-candidatures-yasin-boukraa.xlsx")
STATUTS = ["À envoyer", "À vérifier", "Envoyée", "Relancée", "En cours", "Entretien", "Refus", "En pause"]
F = "Arial"
wb = Workbook(); ws = wb.active; ws.title = "Suivi"
thin = Side(style="thin", color="BFBFBF"); B = Border(left=thin, right=thin, top=thin, bottom=thin)
ws["A1"] = "Candidatures de stage 2027 - Yasin Boukraa"; ws["A1"].font = Font(name=F, size=14, bold=True)
ws["A2"] = f"Mis à jour le {d['maj']} · stage de 2 mois minimum, d'avril à fin août 2027"; ws["A2"].font = Font(name=F, size=10, italic=True, color="595959")
H = ["#", "Entreprise", "Ville", "Type", "CV", "Mode", "Contact", "Envoyée le", "Statut", "Dernière réponse", "Réponse le", "Prochaine action", "Page candidature"]
hr = 4
for i, h in enumerate(H, 1):
    c = ws.cell(hr, i, h); c.font = Font(name=F, bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F3864")
    c.alignment = Alignment(vertical="center", wrap_text=True); c.border = B
def dt(s):
    return datetime.datetime.strptime(s, "%d/%m/%Y").date() if s else None
for r, L in enumerate(d["lignes"], hr + 1):
    name, ville, typ, cv, mode, contact, env, st, rep, repd, action, slug = L
    vals = [r - hr, name, ville, typ, cv, mode, contact, dt(env), st, rep, dt(repd), action,
            f"https://saykito.github.io/Portfolio/candidature/{slug}/" if slug else ""]
    for i, v in enumerate(vals, 1):
        c = ws.cell(r, i, v); c.font = Font(name=F, size=10); c.border = B
        c.alignment = Alignment(vertical="top", wrap_text=True)
        if i in (8, 11): c.number_format = "DD/MM/YYYY"
    if slug:
        c = ws.cell(r, 13); c.hyperlink = c.value; c.font = Font(name=F, size=10, color="0563C1", underline="single")
last = hr + len(d["lignes"])
dv = DataValidation(type="list", formula1='"' + ",".join(STATUTS) + '"', allow_blank=True)
ws.add_data_validation(dv); dv.add(f"I{hr+1}:I{last+30}")
colors = {"Refus": "F8D7DA", "En cours": "D4EDDA", "Entretien": "C3E6CB", "Envoyée": "DCE9F7", "Relancée": "E7E0F7", "À vérifier": "FFF3CD", "À envoyer": "FFF3CD", "En pause": "E7E6E6"}
for k, col in colors.items():
    ws.conditional_formatting.add(f"I{hr+1}:I{last+30}", CellIsRule(operator="equal", formula=[f'"{k}"'], fill=PatternFill("solid", fgColor=col)))
for col, w in zip("ABCDEFGHIJKLM", [4, 30, 22, 16, 11, 11, 34, 12, 12, 34, 12, 34, 44]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = f"C{hr+1}"; ws.auto_filter.ref = f"A{hr}:M{last}"
# Bilan
b = wb.create_sheet("Bilan")
b["A1"] = "Bilan des candidatures"; b["A1"].font = Font(name=F, size=14, bold=True)
b["A3"] = "Statut"; b["B3"] = "Nombre"
for c in (b["A3"], b["B3"]): c.font = Font(name=F, bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F3864")
rng = f"Suivi!$I${hr+1}:$I${last+30}"
for i, s in enumerate(STATUTS, 4):
    b.cell(i, 1, s).font = Font(name=F); b.cell(i, 2, f'=COUNTIF({rng},A{i})').font = Font(name=F)
t = 4 + len(STATUTS)
b.cell(t, 1, "Total").font = Font(name=F, bold=True); b.cell(t, 2, f"=SUM(B4:B{t-1})").font = Font(name=F, bold=True)
b.cell(t + 2, 1, "Mails ou formulaires partis").font = Font(name=F)
b.cell(t + 2, 2, f'=COUNTA(Suivi!$H${hr+1}:$H${last+30})').font = Font(name=F)
b.cell(t + 3, 1, "Réponses reçues (hors accusés auto)").font = Font(name=F)
b.cell(t + 3, 2, f'=COUNTIF({rng},"En cours")+COUNTIF({rng},"Entretien")+COUNTIF({rng},"Refus")').font = Font(name=F)
b.cell(t + 5, 1, "Le statut se choisit dans la liste déroulante de la colonne I de l'onglet Suivi ; ce tableau se met à jour tout seul.").font = Font(name=F, size=9, italic=True, color="595959")
b.column_dimensions["A"].width = 38; b.column_dimensions["B"].width = 10
wb.save(OUT); print(OUT)

# Export CSV pour Google Sheets (upload texte fiable via le connecteur Drive)
import csv
CSV = os.path.join(ROOT, "_outils", "suivi-candidatures-yasin-boukraa.csv")
with open(CSV, "w", encoding="utf-8", newline="") as f:
    w = csv.writer(f)
    n = len(d["lignes"])
    w.writerow(H + ["", "Bilan", "Nombre"])
    sts = [L[7] for L in d["lignes"]]
    bil = [(s, sts.count(s)) for s in STATUTS]
    bil += [("Total", n),
            ("Mails ou formulaires partis", sum(1 for L in d["lignes"] if L[6])),
            ("Réponses reçues (hors accusés auto)", sum(sts.count(s) for s in ("En cours", "Entretien", "Refus"))),
            ("Mis à jour le", d["maj"])]
    for i, L in enumerate(d["lignes"]):
        name, ville, typ, cv, mode, contact, env, st, rep, repd, action, slug = L
        row = [i + 1, name, ville, typ, cv, mode, contact, env, st, rep, repd, action,
               f"https://saykito.github.io/Portfolio/candidature/{slug}/" if slug else ""]
        row += ["", bil[i][0], bil[i][1]] if i < len(bil) else []
        w.writerow(row)
print(CSV)
